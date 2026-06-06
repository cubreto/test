#!/usr/bin/env python3
"""
Generate the technical deck for the Code Graph Modernization sandbox.

    python deck/build_deck.py            # -> deck/code-graph-modernization.pptx

Editable PowerPoint. Re-run to regenerate from this single source of truth.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ---- palette -------------------------------------------------------------
NAVY   = RGBColor(0x0B, 0x24, 0x47)
TEAL   = RGBColor(0x14, 0x9E, 0xC8)
AMBER  = RGBColor(0xE8, 0x9B, 0x1D)
SLATE  = RGBColor(0x3A, 0x4A, 0x5A)
LIGHT  = RGBColor(0xF2, 0xF5, 0xF8)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
GREEN  = RGBColor(0x2E, 0x8B, 0x57)
MUTED  = RGBColor(0x8A, 0x97, 0xA3)

EMU_W, EMU_H = Inches(13.333), Inches(7.5)
MONO = "Consolas"
SANS = "Calibri"

prs = Presentation()
prs.slide_width  = EMU_W
prs.slide_height = EMU_H
BLANK = prs.slide_layouts[6]


def _set(run, size, color, bold=False, font=SANS, italic=False):
    run.font.size = Pt(size); run.font.color.rgb = color
    run.font.bold = bold; run.font.name = font; run.font.italic = italic


def box(slide, x, y, w, h, fill=None, line=None, line_w=1.0):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(line_w)
    return sh


def text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=6, line_spacing=1.0):
    """runs: list of paragraphs; each paragraph is list of (txt,size,color,bold,font,italic)."""
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space_after); p.line_spacing = line_spacing
        for seg in para:
            txt, size, color = seg[0], seg[1], seg[2]
            bold = seg[3] if len(seg) > 3 else False
            font = seg[4] if len(seg) > 4 else SANS
            italic = seg[5] if len(seg) > 5 else False
            r = p.add_run(); r.text = txt; _set(r, size, color, bold, font, italic)
    return tb


def header(slide, kicker, title):
    box(slide, 0, 0, EMU_W, Inches(1.25), fill=WHITE)
    box(slide, 0, Inches(1.18), EMU_W, Pt(3), fill=TEAL)
    box(slide, Inches(0.0), 0, Pt(10), Inches(1.25), fill=TEAL)
    text(slide, Inches(0.55), Inches(0.18), Inches(12), Inches(0.4),
         [[(kicker.upper(), 12, TEAL, True)]])
    text(slide, Inches(0.5), Inches(0.48), Inches(12.4), Inches(0.7),
         [[(title, 26, NAVY, True)]])


def content_slide(kicker, title):
    s = prs.slides.add_slide(BLANK)
    box(s, 0, 0, EMU_W, EMU_H, fill=LIGHT)
    header(s, kicker, title)
    return s


def bullets(slide, items, x=Inches(0.6), y=Inches(1.5), w=Inches(12.1), h=Inches(5.6),
            size=16, gap=10):
    """items: list of (text, level, color?) ; level 0 = •, level1 = sub –"""
    paras = []
    for it in items:
        t, lvl = it[0], it[1]
        color = it[2] if len(it) > 2 else SLATE
        bold = it[3] if len(it) > 3 else False
        if lvl == 0:
            paras.append([("▸  ", size, TEAL, True), (t, size, color, bold)])
        elif lvl == -1:  # plain heading line, amber
            paras.append([(t, size, color, True)])
        else:
            paras.append([("      –  ", size-2, MUTED), (t, size-2, color, bold)])
    text(slide, x, y, w, h, paras, space_after=gap, line_spacing=1.05)


def chip(slide, x, y, w, label, fill, txtcolor=WHITE, size=12, h=Inches(0.38)):
    b = box(slide, x, y, w, h, fill=fill)
    tf = b.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label; _set(r, size, txtcolor, True)
    return b


def table(slide, rows, x, y, w, col_w, header_fill=NAVY, size=12, row_h=Inches(0.42)):
    nrows, ncols = len(rows), len(rows[0])
    h = row_h * nrows
    gt = slide.shapes.add_table(nrows, ncols, x, y, w, h).table
    for j, cw in enumerate(col_w):
        gt.columns[j].width = cw
    for i, row in enumerate(rows):
        gt.rows[i].height = row_h
        for j, val in enumerate(row):
            c = gt.cell(i, j)
            c.margin_left = Inches(0.08); c.margin_right = Inches(0.06)
            c.margin_top = Inches(0.02); c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run(); r.text = str(val)
            if i == 0:
                c.fill.solid(); c.fill.fore_color.rgb = header_fill
                _set(r, size, WHITE, True)
            else:
                c.fill.solid(); c.fill.fore_color.rgb = WHITE if i % 2 else LIGHT
                mono = j == 0
                _set(r, size-0.5, SLATE, False, MONO if mono else SANS)
    return gt


# =========================================================================
# 1. TITLE
# =========================================================================
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, EMU_W, EMU_H, fill=NAVY)
box(s, 0, Inches(4.55), EMU_W, Pt(4), fill=TEAL)
box(s, Inches(0.9), Inches(2.0), Inches(1.6), Pt(6), fill=AMBER)
text(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(2.2),
     [[("Code Graph Modernization", 44, WHITE, True)],
      [("A trustworthy code knowledge graph for legacy modernization", 22, TEAL, False)]],
     space_after=10)
text(s, Inches(0.9), Inches(4.8), Inches(11.5), Inches(1.6),
     [[("Deterministic graph extraction · KDM-lite schema · GraphRAG & RAG on top", 16, LIGHT)],
      [("Open-source learning sandbox  →  TCS / MasterCraft transition", 16, MUTED)]],
     space_after=8)
text(s, Inches(0.9), Inches(6.7), Inches(11.5), Inches(0.5),
     [[("Technical review for Brajesh", 14, WHITE, True),
       ("      ·   June 2026", 14, MUTED)]])

# =========================================================================
# 2. THE PROBLEM
# =========================================================================
s = content_slide("Why this exists", "Modernization needs a graph you can trust")
bullets(s, [
    ("Legacy estates (COBOL / JCL / copybooks) have the business logic but no reliable map: what calls what, what reads/writes which data, where the integration seams are.", 0),
    ("Modernization decisions — sequencing, risk, scope — hinge on that map. Getting it wrong is expensive.", 0),
    ("The tempting shortcut: point an LLM at raw code and ask it to describe the system.", 0, SLATE, True),
    ("Plausible prose, but non-deterministic and unverifiable — you cannot stake a migration plan on it.", 1, RGBColor(0xB0,0x3A,0x2E)),
    ("What we actually need: a typed, evidence-backed dependency graph, with LLM techniques layered on top to explain and draft — never to invent the structure.", 0, NAVY, True),
], y=Inches(1.7), gap=14)
box(s, Inches(0.6), Inches(6.35), Inches(12.1), Inches(0.7), fill=NAVY)
text(s, Inches(0.8), Inches(6.45), Inches(11.7), Inches(0.5),
     [[("Design principle:  ", 15, AMBER, True),
       ("structure is derived deterministically; GraphRAG/LLM sits on top; every claim cites source evidence.", 15, WHITE)]])

# =========================================================================
# 3. ARCHITECTURE / PIPELINE
# =========================================================================
s = content_slide("Architecture", "One pipeline, four stages, evidence end-to-end")
stages = [
    ("Legacy source", "COBOL · copybooks · JCL · SQL", SLATE),
    ("[1] Extractor", "heuristic now →\nMasterCraft adapter later", TEAL),
    ("[2] KDM-lite graph", "typed nodes + edges\n+ source-line evidence", NAVY),
    ("[3] Cluster", "functional-area\ndetection", TEAL),
    ("[4] Grounded specs", "evidence-cited\nmarkdown specs", NAVY),
]
x = Inches(0.55); y = Inches(2.1); bw = Inches(2.28); bh = Inches(1.7); gap = Inches(0.18)
for i, (t, sub, col) in enumerate(stages):
    bx = x + i * (bw + gap)
    box(s, bx, y, bw, bh, fill=WHITE, line=col, line_w=2.0)
    box(s, bx, y, bw, Pt(7), fill=col)
    text(s, bx + Inches(0.12), y + Inches(0.18), bw - Inches(0.24), Inches(0.6),
         [[(t, 14, col, True)]], align=PP_ALIGN.CENTER)
    text(s, bx + Inches(0.12), y + Inches(0.72), bw - Inches(0.24), Inches(0.9),
         [[(line, 11, SLATE)] for line in sub.split("\n")],
         align=PP_ALIGN.CENTER, space_after=2)
    if i < len(stages) - 1:
        text(s, bx + bw - Inches(0.02), y + Inches(0.6), Inches(0.22), Inches(0.5),
             [[("▶", 16, AMBER, True)]], align=PP_ALIGN.CENTER)
bullets(s, [
    ("Branch at [2]: load into Neo4j for interactive Cypher / GDS, and run GraphRAG + simple-RAG over the same corpus.", 0),
    ("Core path is stdlib-only — no DB, no pip — so it runs anywhere, including locked-down sandboxes.", 0, GREEN, True),
], y=Inches(4.35), gap=10)
box(s, Inches(0.6), Inches(6.4), Inches(12.1), Inches(0.6), fill=LIGHT, line=TEAL)
text(s, Inches(0.8), Inches(6.48), Inches(11.7), Inches(0.45),
     [[("Output  →  functional specs + integration points  →  modernization backlog (next phase)", 14, NAVY, True)]])

# =========================================================================
# 4. KDM-LITE SCHEMA
# =========================================================================
s = content_slide("Schema", "KDM-lite: an open stand-in for the MasterCraft inventory")
text(s, Inches(0.6), Inches(1.4), Inches(12), Inches(0.5),
     [[("A property-graph subset of OMG KDM (ISO/IEC 19506), the standard MasterCraft aligns to. "
        "Same node/edge concepts MasterCraft emits — so only the extractor swaps at the TCS boundary.", 14, SLATE)]])
table(s, [
    ["KDM-lite node", "KDM origin", "MasterCraft"],
    ["Program",   "code:CompilationUnit", "Program / Class"],
    ["Function",  "code:CallableUnit",    "Paragraph / Method"],
    ["Screen",    "ui:Screen",            "Screen / Map"],
    ["DataStore", "data:RelationalTable", "File / Table / Queue"],
    ["Job",       "platform:DeployedComp.", "Batch Job / JCL Step"],
], Inches(0.6), Inches(2.25), Inches(6.0),
   [Inches(1.9), Inches(2.2), Inches(1.9)], row_h=Inches(0.46))
table(s, [
    ["KDM-lite edge", "MasterCraft"],
    ["CALLS",            "CALLS"],
    ["READS_FROM",       "READS_FROM"],
    ["WRITES_TO",        "WRITES_TO"],
    ["USES_COPYBOOK",    "USES_COPYBOOK"],
    ["JOB_RUNS_PROGRAM", "JOB_RUNS_PROGRAM"],
], Inches(6.85), Inches(2.25), Inches(5.85),
   [Inches(3.0), Inches(2.85)], row_h=Inches(0.46))
box(s, Inches(0.6), Inches(5.9), Inches(12.1), Inches(0.95), fill=NAVY)
text(s, Inches(0.8), Inches(6.0), Inches(11.7), Inches(0.8),
     [[("Evidence contract (non-negotiable):  ", 14, AMBER, True),
       ("every node & edge carries source_artifact, source_lines, extraction_method, confidence.", 14, WHITE)],
      [("Deterministic facts are distinguishable from inferred ones; sandbox stamps heuristic-cobol-v0 / 0.7, MasterCraft would stamp ≈1.0.", 12, MUTED)]],
     space_after=6)

# =========================================================================
# 5. EXTRACTOR + EVIDENCE
# =========================================================================
s = content_slide("The extractor", "Deterministic structure, every fact stamped")
bullets(s, [
    ("Heuristic COBOL/JCL reader (stdlib only) → typed KDM-lite nodes & edges.", 0),
    ("Detects: program CALLS, file/table READS_FROM / WRITES_TO, USES_COPYBOOK, screen → program, JCL job → program.", 1),
    ("Every fact carries its provenance — nothing is asserted without a source line:", 0, NAVY, True),
], y=Inches(1.55), gap=10)
box(s, Inches(0.6), Inches(3.05), Inches(12.1), Inches(1.6), fill=RGBColor(0x10,0x1A,0x28))
text(s, Inches(0.85), Inches(3.2), Inches(11.6), Inches(1.4),
     [[("ACCTPOST", 13, TEAL, True, MONO), ("  --CALLS-->  ", 13, MUTED, False, MONO),
       ("(missing)        ", 13, WHITE, False, MONO)],
      [("CUSTMGMT", 13, TEAL, True, MONO), ("  --READS_FROM-->  ", 13, MUTED, False, MONO),
       ("CUSTOMER-FILE", 13, WHITE, True, MONO)],
      [("   source_artifact: ", 12, MUTED, False, MONO), ("CUSTMGMT.CBL", 12, AMBER, False, MONO),
       ("   lines: ", 12, MUTED, False, MONO), ("28", 12, AMBER, False, MONO),
       ("   method: ", 12, MUTED, False, MONO), ("heuristic-cobol-v0", 12, GREEN, False, MONO),
       ("   confidence: ", 12, MUTED, False, MONO), ("0.7", 12, GREEN, False, MONO)]],
     space_after=8, line_spacing=1.1)
bullets(s, [
    ("Low-confidence facts (e.g. inferred screens, 0.5) are queryable separately — reviewers audit them first.", 0),
    ("At the TCS boundary the heuristic is replaced by a MasterCraft-inventory adapter; same graph, facts stamped deterministic (~1.0).", 0, NAVY, True),
], y=Inches(4.95), gap=12)

# =========================================================================
# 6. SAMPLE RESULTS
# =========================================================================
s = content_slide("Proof on a sample estate", "16 nodes · 19 edges · 3 functional areas")
metrics = [("4", "programs"), ("3", "copybooks"), ("1", "JCL job"),
           ("2", "SQL tables"), ("3", "functional areas")]
mx = Inches(0.6); mw = Inches(2.3); my = Inches(1.55)
for i, (n, lbl) in enumerate(metrics):
    bx = mx + i * (mw + Inches(0.13))
    box(s, bx, my, mw, Inches(1.15), fill=NAVY)
    box(s, bx, my, Pt(6), Inches(1.15), fill=TEAL)
    text(s, bx, my + Inches(0.12), mw, Inches(0.6), [[(n, 30, WHITE, True)]], align=PP_ALIGN.CENTER)
    text(s, bx, my + Inches(0.74), mw, Inches(0.35), [[(lbl, 13, TEAL)]], align=PP_ALIGN.CENTER)
bullets(s, [
    ("Customer mgmt & daily reporting — CUSTMGMT (screen), DAILYRPT (batch via DAILYJOB) over CUSTOMER-FILE / REPORT-FILE.", 0),
    ("Account posting — ACCTPOST over ACCOUNT / TRANSACTION (DB2).", 0),
    ("Parts ordering — PARTORDR (screen) over PARTS / ORDERS.", 0),
], y=Inches(3.0), gap=9, size=15)
box(s, Inches(0.6), Inches(4.75), Inches(12.1), Inches(2.05), fill=LIGHT, line=AMBER, line_w=2.0)
text(s, Inches(0.85), Inches(4.9), Inches(11.6), Inches(1.85),
     [[("Cross-area edges = the integration contracts that drive sequencing", 16, NAVY, True)],
      [("CUSTMGMT  CALLS  ACCTPOST", 14, SLATE, True, MONO),
       ("   — customer screen reaches into account posting", 13, SLATE)],
      [("DAILYRPT  READS_FROM  ACCOUNT", 14, SLATE, True, MONO),
       ("   — reporting depends on the posting area's data", 13, SLATE)],
      [("These cross-cluster dependencies are the contracts and risks to manage when you carve the estate.", 13, SLATE, False, SANS, True)]],
     space_after=8, line_spacing=1.05)

# =========================================================================
# 7. THREE AI PATHS
# =========================================================================
s = content_slide("Three retrieval/generation paths", "Same corpus, different trust profiles")
table(s, [
    ["Path", "How it works", "Output", "Trust"],
    ["Graph + spec gen", "Deterministic parse → typed graph → cluster → spec", "Evidence-cited specs", "High / auditable"],
    ["GraphRAG", "LLM extracts entities+rels → communities → reports", "Discovery, summaries", "Needs validation"],
    ["Simple RAG", "Chunk → retrieve top-k → stuff → LLM answer", "Q&A over corpus", "Baseline"],
], Inches(0.6), Inches(1.55), Inches(12.1),
   [Inches(2.4), Inches(4.6), Inches(2.7), Inches(2.4)], row_h=Inches(0.85), size=13)
bullets(s, [
    ("Primary path is the deterministic graph — structure you can stake a plan on.", 0, NAVY, True),
    ("GraphRAG and RAG are complementary: great for discovery and natural-language exploration, validated against the graph.", 0),
    ("Reliability vs coverage: the deterministic path gives typed, stamped edges; GraphRAG gives breadth but untyped, inferred edges.", 0, SLATE),
], y=Inches(5.1), gap=10, size=14)

# =========================================================================
# 8. GRAPHRAG PATH DETAIL
# =========================================================================
s = content_slide("GraphRAG path", "Microsoft GraphRAG 3.1, tuned for code")
bullets(s, [
    ("Code-tuned config: entity_types set to the KDM-lite vocabulary (program, function, screen, copybook, datastore, job, handler).", 0),
    ("Extraction prompt rewritten with COBOL examples → emits CALLS / READS_FROM / WRITES_TO / USES_COPYBOOK / JOB_RUNS_PROGRAM.", 0),
    ("GraphML snapshot on → load the LLM-built graph into Neo4j alongside the deterministic one.", 0),
    ("Runs two ways:", 0, NAVY, True),
    ("In Claude-on-the-web sandbox: Anthropic completion, embeddings-free workflow (extract → communities → reports → global search).", 1),
    ("Locally: OpenAI / Together with embeddings → full pipeline incl. local & DRIFT search.", 1),
    ("Verified end-to-end to the LLM boundary: loads docs → chunks → dispatches a Claude extraction call per document. A valid key completes the run.", 0, GREEN, True),
], y=Inches(1.5), gap=8, size=14)
box(s, Inches(0.6), Inches(6.5), Inches(12.1), Inches(0.6), fill=NAVY)
text(s, Inches(0.8), Inches(6.58), Inches(11.7), Inches(0.45),
     [[("Network reality:  ", 13, AMBER, True),
       ("sandbox allowlists Anthropic only; OpenAI/Together hosts run locally. Config documents both paths.", 13, WHITE)]])

# =========================================================================
# 9. TCS / MASTERCRAFT TRANSITION
# =========================================================================
s = content_slide("Transition to TCS / MasterCraft", "One box changes — everything downstream is reused")
y = Inches(1.7)
seq = [("MasterCraft\ninventory export", AMBER, True),
       ("Adapter\n(maps to KDM-lite)", TEAL, True),
       ("Same graph.json", NAVY, False),
       ("Neo4j · cluster ·\nspecs · GraphRAG", SLATE, False)]
bw = Inches(2.75); gap = Inches(0.35); x = Inches(0.7)
for i, (t, col, changes) in enumerate(seq):
    bx = x + i * (bw + gap)
    box(s, bx, y, bw, Inches(1.5), fill=col if changes else WHITE, line=col, line_w=2.5)
    tcol = WHITE if changes else col
    text(s, bx + Inches(0.1), y + Inches(0.35), bw - Inches(0.2), Inches(0.9),
         [[(line, 14, tcol, True)] for line in t.split("\n")],
         align=PP_ALIGN.CENTER, space_after=2)
    if changes:
        chip(s, bx + Inches(0.5), y - Inches(0.42), bw - Inches(1.0), "CHANGES", AMBER, NAVY, 10)
    else:
        chip(s, bx + Inches(0.55), y - Inches(0.42), bw - Inches(1.1), "REUSED", GREEN, WHITE, 10)
    if i < len(seq) - 1:
        text(s, bx + bw, y + Inches(0.5), gap, Inches(0.5), [[("▶", 18, SLATE, True)]], align=PP_ALIGN.CENTER)
bullets(s, [
    ("Choosing a KDM-aligned schema for the sandbox is what makes this a remapping of sources, not a redesign.", 0, NAVY, True),
    ("graph.json format, Neo4j loader, clustering, spec generation, GraphRAG config — all reused unchanged.", 0),
    ("De-risks the TCS engagement: the hard parts (schema, clustering, evidence, specs) are proven on open code first.", 0, GREEN, True),
], y=Inches(3.9), gap=12, size=15)

# =========================================================================
# 10. ROADMAP
# =========================================================================
s = content_slide("Roadmap", "Where we are and what's next")
phases = [
    ("0", "Schema alignment", "Confirm MasterCraft export maps onto KDM-lite nodes/edges", "Ready"),
    ("1", "This sandbox", "Prove extract → graph → cluster → spec on open COBOL", "Done"),
    ("2", "GraphRAG enrichment", "Semantic retrieval + richer narratives; broaden coverage", "In progress"),
    ("3", "Target mapping", "MAPS_TO_TARGET edges, capability specs, modernization backlog", "Planned"),
]
status_col = {"Done": GREEN, "In progress": AMBER, "Ready": TEAL, "Planned": MUTED}
y = Inches(1.65); rh = Inches(1.18)
for i, (n, title, desc, st) in enumerate(phases):
    by = y + i * (rh + Inches(0.08))
    box(s, Inches(0.6), by, Inches(12.1), rh, fill=WHITE, line=RGBColor(0xD5,0xDD,0xE3))
    box(s, Inches(0.6), by, Inches(1.0), rh, fill=NAVY)
    text(s, Inches(0.6), by + Inches(0.32), Inches(1.0), Inches(0.6), [[(n, 26, WHITE, True)]], align=PP_ALIGN.CENTER)
    text(s, Inches(1.8), by + Inches(0.16), Inches(7.6), Inches(0.9),
         [[(title, 17, NAVY, True)], [(desc, 13, SLATE)]], space_after=3)
    chip(s, Inches(10.65), by + Inches(0.38), Inches(1.85), st, status_col[st], WHITE, 12, Inches(0.42))

# =========================================================================
# 11. CLOSING / NEXT STEPS
# =========================================================================
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, EMU_W, EMU_H, fill=NAVY)
box(s, Inches(0.9), Inches(0.85), Inches(1.6), Pt(6), fill=AMBER)
text(s, Inches(0.9), Inches(1.05), Inches(11.5), Inches(0.8), [[("What this gives us", 32, WHITE, True)]])
bullets_n = [
    ("A proven, evidence-first pipeline from legacy source to grounded functional specs.", GREEN),
    ("A KDM-aligned graph that transfers to TCS/MasterCraft by swapping one component.", TEAL),
    ("GraphRAG + RAG layered on top for discovery — validated against a deterministic spine.", TEAL),
    ("De-risked path: the architecture is exercised on open COBOL before touching client code.", GREEN),
]
paras = [[("▸  ", 18, AMBER, True), (t, 18, WHITE)] for (t, c) in bullets_n]
text(s, Inches(0.95), Inches(2.0), Inches(11.6), Inches(2.8), paras, space_after=16, line_spacing=1.1)
box(s, Inches(0.9), Inches(4.75), Inches(11.5), Inches(2.0), fill=RGBColor(0x10,0x33,0x55), line=TEAL, line_w=1.5)
text(s, Inches(1.2), Inches(4.95), Inches(11.0), Inches(1.8),
     [[("Asks / next steps", 18, AMBER, True)],
      [("1.  Get a MasterCraft inventory export sample to validate the KDM-lite mapping (Phase 0).", 15, WHITE)],
      [("2.  Confirm target-state vocabulary so we can design MAPS_TO_TARGET edges (Phase 3).", 15, WHITE)],
      [("3.  Decide GraphRAG scope: discovery aid vs. part of the deliverable.", 15, WHITE)]],
     space_after=8, line_spacing=1.05)

# -------------------------------------------------------------------------
out = os.path.join(os.path.dirname(__file__), "code-graph-modernization.pptx")
prs.save(out)
print(f"wrote {out}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
